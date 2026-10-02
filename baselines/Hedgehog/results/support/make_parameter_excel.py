from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


OUT = Path(__file__).resolve().parents[1] / "hedgehog_parameter_comparison.xlsx"


def style_sheet(sheet):
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="C8C8C8")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(wrap_text=True, vertical="top")

    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row in sheet.iter_rows():
        for cell in row:
            cell.border = border
            cell.alignment = wrap

    sheet.freeze_panes = "A2"
    for col in sheet.columns:
        max_len = 0
        for cell in col:
            value = "" if cell.value is None else str(cell.value)
            max_len = max(max_len, min(len(value), 60))
        sheet.column_dimensions[get_column_letter(col[0].column)].width = max(14, min(max_len + 4, 62))


def add_sheet(workbook, title, headers, rows):
    sheet = workbook.create_sheet(title)
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    style_sheet(sheet)
    return sheet


def main():
    workbook = Workbook()
    overview = workbook.active
    overview.title = "Overview"
    overview.append(["类别", "结论"])
    overview.append(["论文实际给出", "主要给出几何约束、流程原则、实时系统配置；很少给优化数值。"])
    overview.append(["当前实现补充", "为离线复现补充了相机、固定视角、迭代次数、force 权重、步长、pole 长度范围和 evaluator 指标。"])
    overview.append(["复现性质", "不是论文原始源码/原参数复现，而是按论文约束实现的离线启发式复现。"])
    style_sheet(overview)
    overview.column_dimensions["A"].width = 18
    overview.column_dimensions["B"].width = 80

    add_sheet(
        workbook,
        "Paper_Params",
        ["项目", "论文中的值/说法", "依据位置", "性质"],
        [
            ["运行帧率", "约 30Hz", "论文 Implementation 部分明确说明", "实际参数"],
            ["渲染分辨率", "640 x 480", "论文 Implementation 部分明确说明", "实际参数"],
            ["实现语言/框架", "C++ / OpenSceneGraph", "论文 Implementation 部分明确说明", "实际参数"],
            ["硬件", "Windows 7, Intel i7 quad-core 2.66GHz, 12GB RAM, Nvidia 780GTX", "论文 Implementation 部分明确说明", "实际参数"],
            ["AR 输入", "KinFu / KinectFusion + Microsoft Xbox360 Kinect", "论文 Implementation 部分明确说明", "实际参数"],
            ["pole 方向", "从物体 bounding sphere center 指向 anchor point 的径向方向", "论文 Method 3.1", "算法约束"],
            ["1D Hedgehog", "annotation 只能沿 3D pole 移动，即只改变 pole length", "论文 Method 3.2", "算法约束"],
            ["3D Hedgehog", "沿 pole 移动 + annotation local X/Y plane 内移动", "论文 Method 3.2", "算法约束"],
            ["3D 平面内移动限制", "限制在 annotation size 内", "论文 Method 3.2", "算法约束"],
            ["Plane 方向", "与 viewing plane 平行，或用户定义在 world space", "论文 Method 3.2 / Conclusion", "算法约束"],
            ["Plane 数量", "由用户运行时设置；Figure 7 示例为 3 个", "论文 Plane-Based / Future Work", "半参数：示例给 3，但不是固定默认值"],
            ["Plane 分配", "label 分配到离 anchor point 最近的 plane", "论文 Plane-Based", "算法约束"],
            ["Plane 更新策略", "布局生成后冻结；view vector 与 plane normal 夹角超过用户阈值后再更新", "论文 Plane-Based", "算法约束，但阈值未给"],
            ["优化方法", "force-based optimization；Plane-Based 提到 Ali et al. spring embedding", "论文 Method 3.2", "方法原则，非数值参数"],
            ["角度阈值数值", "未公开", "论文只说 user-defined threshold", "缺失"],
            ["force 权重", "未公开", "论文未列出", "缺失"],
            ["迭代次数/步长", "未公开", "论文未列出", "缺失"],
            ["PCK/OLR/LCD", "论文没有这些量化指标", "当前实现新增 evaluator", "非论文参数"],
        ],
    )

    add_sheet(
        workbook,
        "Impl_Params",
        ["参数", "当前实现值", "用途", "与论文关系"],
        [
            ["CAMERA_RADIUS", "10.0", "离线预览/评估相机半径", "论文未给；当前实现自定"],
            ["FOCAL_LENGTH_MM", "50.0", "离线投影焦距", "论文未给；当前实现自定"],
            ["SENSOR_WIDTH_MM", "36.0", "相机传感器宽度", "论文未给；当前实现自定"],
            ["SENSOR_HEIGHT_MM", "24.0", "相机传感器高度", "论文未给；当前实现自定"],
            ["CAMERA_NEAR", "1e-6", "近裁剪面", "论文未给；当前实现自定"],
            ["CAMERA_FAR", "1e6", "远裁剪面", "论文未给；当前实现自定"],
            ["PERTURB_DEGREES", "45.0", "生成 up/down/left/right 视角扰动角", "论文未给；当前实现自定"],
            ["PREVIEW_WIDTH", "750", "当前 preview/evaluator 宽度", "论文实际实现为 640；当前实现自定"],
            ["PREVIEW_HEIGHT", "500", "当前 preview/evaluator 高度", "论文实际实现为 480；当前实现自定"],
            ["iterations", "默认 80", "优化迭代上限", "论文未给；当前实现自定"],
            ["plane_count", "默认 3", "平面数量", "论文示例用 3，但论文说用户运行时设置；当前默认取 3"],
            ["min_len", "0.08 * object_diagonal", "Hedgehog pole 最小长度", "论文未给；当前实现自定"],
            ["max_len", "2.2 * object_diagonal + max_label_size", "Hedgehog pole 最大长度", "论文未给；当前实现自定"],
            ["收敛阈值", "1e-7", "force norm 小于该值提前停止", "论文未给；当前实现自定"],
            ["Hedgehog force 缩放", "0.5 * force", "屏幕力转世界位移前缩放", "论文未给；当前实现自定"],
        ],
    )

    add_sheet(
        workbook,
        "Force_Params",
        ["参数/权重", "当前实现值", "用途", "与论文关系"],
        [
            ["label-label separation weight", "0.65", "标签框重叠时互相推开", "论文只说 force-based，未给权重"],
            ["object overlap weight", "95.0 * overlap / area", "标签覆盖物体投影 bbox 时向外推", "论文未给；且当前用 bbox 近似真实遮挡"],
            ["object overlap cap", "180.0", "限制物体重叠推力上限", "论文未给"],
            ["anchor occlusion force", "26.0", "标签遮住其他 anchor 时推开", "论文未给"],
            ["leader crossing force", "18.0", "引导线交叉时推开", "论文未给"],
            ["compactness pull-back", "0.012 * (center - anchor)", "防止标签离 anchor 过远", "论文未给"],
            ["force clip", "[-80.0, 80.0]", "限制单轮屏幕力范围", "论文未给"],
        ],
    )

    add_sheet(
        workbook,
        "Plane_Params",
        ["项目", "当前实现值", "用途", "与论文关系"],
        [
            ["plane depth sampling", "camera depth min/max 等距采样", "生成 view-parallel planes", "近似论文 screen-aligned bounding box 内等距 plane"],
            ["initial outward distance", "max(35.0, 0.35 * mean_label * pixel_scale)", "初始化 plane 内 label 位置", "论文未给"],
            ["spring", "0.010", "把 label 往 anchor 拉回", "论文提 spring embedding，但未给数值"],
            ["step", "0.55", "每轮更新步长", "论文未给"],
            ["max_offset", "0.85 * max(width, height)", "限制 label 离 anchor 太远", "论文未给"],
            ["freeze layout", "未实现", "原文建议布局冻结并按角度阈值更新", "当前简化：每个 view 独立重新计算"],
            ["angle threshold", "未实现", "控制 plane layout 何时更新", "论文说 user-defined，但未给具体值"],
        ],
    )

    add_sheet(
        workbook,
        "Evaluator_Params",
        ["指标/参数", "当前实现值", "用途", "与论文关系"],
        [
            ["PCK_005", "图像对角线 5%", "预测 label center 接近 manual center 的比例", "论文没有；当前 evaluator 新增"],
            ["PCK_010", "图像对角线 10%", "预测 label center 接近 manual center 的比例", "论文没有；当前 evaluator 新增"],
            ["OLR", "overlap / label area", "标签重叠率", "论文没有这个量化定义；当前新增"],
            ["LCD", "leader crossing 三角面积比例近似", "引导线交叉程度", "论文没有这个量化定义；当前新增"],
            ["avg_leader_length", "leader line 屏幕长度 / 图像对角线", "引导线长度", "论文没有这个量化定义；当前新增"],
            ["quality_score", "100*(0.7*PCK_005 + 0.3*PCK_010) - 25*OLR - 10*LCD", "综合分", "完全是当前实现自定义"],
            ["label thickness", "0.02", "投影 label bounds 时使用", "论文未给；数据 box_size 第三维也常为 0.02"],
            ["object occlusion", "projected object bbox overlap", "物体遮挡近似", "论文关注对象遮挡；当前不是深度/mask 精确遮挡"],
        ],
    )

    add_sheet(
        workbook,
        "Flow_Compare",
        ["模块", "论文原流程", "当前实现", "简化点"],
        [
            ["DataLoader", "实时 3D scene / AR 系统输入", "读取 ../data/Layout/*/*/layout1/Annotation/*.json 和 Obj-O/*.obj", "从实时系统简化为离线 JSON/OBJ 批处理"],
            ["Annotation 表示", "完整 3D annotation，可旋转、可朝向相机", "只保留 center、anchor、box_size", "省略 annotation 姿态和真实 billboard 状态"],
            ["Bounding sphere center", "使用物体 bounding sphere center", "使用 OBJ/anchor 的 AABB center 近似", "几何中心计算被简化"],
            ["1D Hedgehog", "只沿 pole 改变长度", "保留；但补 min/max length、迭代和 force 权重", "约束一致，数值自定"],
            ["3D Hedgehog", "pole + annotation local X/Y plane", "使用 camera x/y 方向近似局部平面，并按 label size clip", "局部坐标和姿态处理被简化"],
            ["Plane-Based", "view-parallel 或 user-defined planes；布局冻结并阈值更新", "固定 view-parallel planes；每个 view 独立重算", "未实现 temporal freeze 和 angle threshold"],
            ["Evaluator", "论文主要 qualitative，未提供该组指标", "PCK/OLR/LCD/leader length/quality score", "当前为离线量化比较新增"],
        ],
    )

    workbook.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
